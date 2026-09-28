"""Import the client's obligation workbook.

    python manage.py import_obligations <path.xlsx> [--dry-run] [--sheet NAME] [--keep-missing]

The workbook needs a header row with the register columns (SPEC 10); the row can be anywhere
in the first 15 rows of the sheet. Obligations missing from the workbook are deactivated
unless --keep-missing is given. The dry run reports what would be created, updated and
deactivated, and writes nothing.
"""

from django.core.management.base import BaseCommand, CommandError
from openpyxl import load_workbook

from apps.obligations.columns import REGISTER_COLUMNS
from apps.obligations.services import sync_obligations


def _norm(v) -> str:
    return " ".join(str(v or "").replace("\n", " ").split()).lower()


def read_workbook(path: str, sheet: str | None = None) -> list[dict]:
    try:
        wb = load_workbook(path, read_only=True, data_only=True)
    except Exception as exc:
        raise CommandError(f"Could not open {path}: {exc}") from exc
    sheets = [wb[sheet]] if sheet else wb.worksheets
    wanted = {_norm(h): key for h, key in REGISTER_COLUMNS}
    for ws in sheets:
        rows = list(ws.iter_rows(values_only=True))
        for idx, row in enumerate(rows[:15]):
            header = {_norm(c): i for i, c in enumerate(row) if c}
            if "obligation id" in header and len(set(header) & set(wanted)) >= 10:
                missing = [h for h, _ in REGISTER_COLUMNS if _norm(h) not in header]
                if missing:
                    raise CommandError(f"Sheet '{ws.title}' lacks columns: {', '.join(missing)}")
                out = []
                for r in rows[idx + 1 :]:
                    rec = {
                        key: r[header[h]] if header[h] < len(r) else None
                        for h, key in wanted.items()
                    }
                    if rec.get("id"):
                        out.append(rec)
                return out
    raise CommandError("No sheet has the register header row (Obligation ID, Regulatory Source…).")


class Command(BaseCommand):
    help = "Import obligations from the client's XLSX workbook."

    def add_arguments(self, parser):
        parser.add_argument("path")
        parser.add_argument("--dry-run", action="store_true")
        parser.add_argument("--sheet", default=None)
        parser.add_argument(
            "--keep-missing",
            action="store_true",
            help="Do not deactivate obligations absent from the workbook.",
        )

    def handle(self, *args, path, dry_run, sheet, keep_missing, **opts):
        rows = read_workbook(path, sheet)
        report = sync_obligations(rows, deactivate_missing=not keep_missing, dry_run=dry_run)
        self.stdout.write(
            ("DRY RUN – nothing written. " if dry_run else "")
            + f"{len(rows)} rows read from {path}."
        )
        for line in report.lines():
            self.stdout.write(line)
