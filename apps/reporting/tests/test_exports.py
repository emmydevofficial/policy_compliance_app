import csv
import io
import zipfile

from openpyxl import load_workbook

from apps.core.tests.helpers import good_submission
from apps.workflow import services as wf
from apps.workflow.files import SUBMISSION_TYPES, validate_upload

CLIENT_WORKBOOK_COLUMNS = [
    "Obligation ID",
    "Regulatory Source",
    "Document Reference",
    "Obligation Category",
    "Obligation Description",
    "Obligation Type",
    "Frequency",
    "Primary Responsible Dept",
    "Supporting Departments",
    "Risk Level",
    "Evidence / Documentation Required",
    "Regulatory Deadline / Trigger",
    "Notes / Penalty for Non-Compliance",
]


def test_register_xlsx_column_order_matches_the_client_workbook_exactly(client, admin, req_a):
    client.force_login(admin)
    r = client.get("/reports/register.xlsx")
    ws = load_workbook(io.BytesIO(r.content)).active
    header = [c.value for c in ws[1]]
    assert header == CLIENT_WORKBOOK_COLUMNS
    assert ws["A2"].value == req_a.obligation.obligation_id


def test_export_writes_an_audit_event(client, executive, req_a):
    from apps.auditlog.models import AuditEvent

    client.force_login(executive)
    client.get("/reports/register.xlsx")
    assert AuditEvent.objects.filter(action="export_generated", actor=executive).exists()


def test_tracker_has_eighteen_columns(client, admin, req_a):
    client.force_login(admin)
    ws = load_workbook(io.BytesIO(client.get("/reports/tracker.xlsx").content)).active
    assert len([c for c in ws[1] if c.value]) == 18


def test_report_pdf_renders(client, admin, req_a):
    client.force_login(admin)
    r = client.get("/reports/report.pdf")
    assert r.status_code == 200 and r.content.startswith(b"%PDF")


def test_submission_pack_contains_verified_files_and_manifest(
    client, admin, officer_a, req_a, pdf_upload
):
    sub = wf.submit(
        req_a,
        officer_a,
        good_submission(),
        [validate_upload(pdf_upload("proof.pdf"), SUBMISSION_TYPES)],
    )
    wf.verify(sub, admin)
    client.force_login(admin)
    r = client.get("/reports/pack.zip")
    zf = zipfile.ZipFile(io.BytesIO(b"".join(r.streaming_content)))
    oid = req_a.obligation.obligation_id
    assert f"{oid}/proof.pdf" in zf.namelist()
    rows = list(csv.DictReader(io.StringIO(zf.read("manifest.csv").decode())))
    att = sub.attachments.get()
    assert rows[0]["sha256"] == att.sha256 and rows[0]["obligation"] == oid
    assert rows[0]["submitted_by"] == officer_a.email


def test_dashboards_show_counts(client, admin, executive, req_a):
    client.force_login(admin)
    assert b"Awaiting verification" in client.get("/dashboard/").content
    client.force_login(executive)
    assert b"Licence exposure" in client.get("/dashboard/executive/").content
